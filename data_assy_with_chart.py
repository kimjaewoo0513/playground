import pandas as pd
import os
from openpyxl import Workbook
from openpyxl.chart import LineChart, Reference

# 추출할 엑셀 파일들이 있는 디렉토리 경로
directory = os.path.normpath(r'C:\Users\floatic-j\Desktop\Dlody 안정성 테스트\240307\원본\새 폴더')

# 결과를 저장할 엑셀 파일 경로 설정
result_file_path = os.path.join(directory, 'result.xlsx')

# 결과를 저장할 DataFrame 생성
result_df = pd.DataFrame()

# 결과 파일이 없는 경우 생성
if not os.path.exists(result_file_path):
    wb = Workbook()
    wb.save(result_file_path)

# 결과를 엑셀 파일로 저장
with pd.ExcelWriter(result_file_path, engine='openpyxl', mode='a') as writer:
    for filename in os.listdir(directory):
        if filename.endswith('.xlsx'):
            df = pd.read_excel(os.path.join(directory, filename), usecols=['timestamp', 'pitch'])

            # 각 파일에서 추출된 데이터를 임시로 저장할 DataFrame 생성
            temp_df = pd.DataFrame({'timestamp': df['timestamp'], 'pitch': df['pitch']})

            # 파일별로 시트를 생성하여 데이터 저장
            temp_sheet_name = os.path.splitext(filename)[0]
            temp_df.to_excel(writer, sheet_name=temp_sheet_name, index=False)
            print(df.columns)

            # 각 파일별로 차트 생성
            worksheet = writer.book.create_sheet(title=f'{temp_sheet_name}_Chart')
            chart = LineChart()
            chart.title = 'Pitch'
            chart.x_axis.title = 'Timestamp'
            chart.y_axis.title = 'Pitch'
            chart.width *= 3
            chart.height *= 1

            # 데이터 범위 설정
            data = Reference(writer.book[temp_sheet_name], min_col=2, min_row=1, max_col=temp_df.shape[1], max_row=temp_df.shape[0])
            categories = Reference(writer.book[temp_sheet_name], min_col=1, min_row=2, max_row=temp_df.shape[0])
            chart.add_data(data, titles_from_data=True)
            chart.set_categories(categories)

            # 차트를 추가할 시트에 추가
            worksheet.add_chart(chart, 'A1')

# 결과 파일 저장
writer.save()
